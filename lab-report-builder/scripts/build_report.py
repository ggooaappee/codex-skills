#!/usr/bin/env python3
"""Rebuild a course lab report from an existing report template.

Preferred invocation (all paths live in a UTF-8 JSON job file, because this
Windows console mangles non-ASCII command-line arguments):

    python build_report.py --job report.json

report.json:
{
  "template": "C:/.../1123...实验二.docx",
  "out":      "C:/.../1123...实验三.docx",
  "date":     "2026.10.12",
  "title":    "STP 与 MSTP 配置",
  "summary":  "本次实验……",
  "row4_min_twips": 630,
  "blocks": [
    {"type": "head", "text": "一、实验目的"},
    {"type": "body", "text": "……"},
    {"type": "cmd",  "lines": ["[SW1]stp mode stp"]},
    {"type": "img",  "path": "C:/.../topo.png", "width": 3.6},
    {"type": "cap",  "text": "图 1  STP 拓扑"}
  ]
}

The script clones the template's own formatting, replaces the body cell, and
updates the date / title / summary rows.
"""

import argparse
import copy
import json
import os

import docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

W_P = qn("w:p")
W_T = qn("w:t")
W_BR = qn("w:br")
W_PPR = qn("w:pPr")
W_RPR = qn("w:rPr")
W_TBL = qn("w:tbl")
W_TR = qn("w:tr")
W_TRPR = qn("w:trPr")
W_TRHEIGHT = qn("w:trHeight")
W_VAL = qn("w:val")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def para_font(para):
    for run in para.runs:
        if run.font.name:
            return run.font.name
    return None


def para_size(para):
    for run in para.runs:
        if run.font.size:
            return run.font.size.pt
    return None


def para_bold(para):
    for run in para.runs:
        if run.bold is not None:
            return run.bold
    return None


def find_templates(cell):
    """Locate formatting prototypes inside the existing report body."""
    templates = {
        "h1": None,
        "head": None,
        "body": None,
        "cmd": None,
        "img": None,
        "cap": None,
    }
    for para in cell.paragraphs:
        if templates["img"] is None and "graphic" in para._p.xml:
            templates["img"] = para
            continue
        font = para_font(para)
        size = para_size(para)
        bold = para_bold(para)
        if font == "Consolas" and templates["cmd"] is None:
            templates["cmd"] = para
        elif font == "黑体" and bold:
            if size and size >= 11:
                if templates["h1"] is None:
                    templates["h1"] = para
            elif templates["head"] is None:
                templates["head"] = para
        elif font == "仿宋" and size == 9 and templates["cap"] is None:
            templates["cap"] = para
        elif font == "仿宋" and not bold and templates["body"] is None:
            templates["body"] = para

    if templates["h1"] is None:
        templates["h1"] = templates["head"]
    missing = [n for n, v in templates.items() if v is None]
    if missing:
        raise SystemExit(
            "无法从模板正文中识别以下段落样式："
            + ", ".join(missing)
            + "\n请确认模板是同一套实验报告（含标题/正文/命令/图注）。"
        )
    return templates


def strip_to_ppr(element):
    for child in list(element):
        if child.tag != W_PPR:
            element.remove(child)


def clone_run_rpr(template_para):
    run = copy.deepcopy(template_para.runs[0]._r)
    for child in list(run):
        if child.tag != W_RPR:
            run.remove(child)
    return run


def mk_text(template_para, text):
    new_p = copy.deepcopy(template_para._p)
    strip_to_ppr(new_p)
    run = clone_run_rpr(template_para)
    node = run.makeelement(W_T, {})
    node.set(XML_SPACE, "preserve")
    node.text = text
    run.append(node)
    new_p.append(run)
    return new_p


def mk_cmd(template_para, lines):
    new_p = copy.deepcopy(template_para._p)
    strip_to_ppr(new_p)
    run = clone_run_rpr(template_para)
    for index, line in enumerate(lines):
        if index:
            run.append(run.makeelement(W_BR, {}))
        node = run.makeelement(W_T, {})
        node.set(XML_SPACE, "preserve")
        node.text = line
        run.append(node)
    new_p.append(run)
    return new_p


def clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag in (W_P, W_TBL):
            tc.remove(child)


def set_cell_text(cell, text):
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph()
        paragraphs = cell.paragraphs
    first = paragraphs[0]
    if not first.runs:
        first.add_run()
    first.runs[0].text = text
    for run in first.runs[1:]:
        run._r.getparent().remove(run._r)
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_row_min_height(table, row_index, twips):
    tr = table._tbl.findall(W_TR)[row_index]
    tr_pr = tr.find(W_TRPR)
    if tr_pr is None:
        tr_pr = tr.makeelement(W_TRPR, {})
        tr.insert(0, tr_pr)
    height = tr_pr.find(W_TRHEIGHT)
    if height is None:
        height = tr_pr.makeelement(W_TRHEIGHT, {})
        tr_pr.append(height)
    height.set(W_VAL, str(twips))
    height.set(qn("w:hRule"), "atLeast")


def build(job):
    template = job["template"]
    out = job["out"]
    document = docx.Document(template)
    table = document.tables[0]
    if len(table.rows) < 6:
        raise SystemExit("模板表格行数不足，请确认是实验报告模板。")

    cell = table.rows[4].cells[1]
    templates = find_templates(cell)
    clear_cell(cell)

    base_dir = os.path.dirname(os.path.abspath(template))
    count = {"h1": 0, "head": 0, "body": 0, "cmd": 0, "img": 0, "cap": 0}
    for block in job.get("blocks", []):
        kind = block.get("type")
        if kind == "img":
            path = block["path"]
            if not os.path.isabs(path):
                path = os.path.join(base_dir, path)
            element = copy.deepcopy(templates["img"]._p)
            strip_to_ppr(element)
            cell._tc.append(element)
            width = float(block.get("width", 4.13))
            Paragraph(element, cell).add_run().add_picture(path, width=Inches(width))
            count["img"] += 1
            continue

        if kind == "cmd":
            element = mk_cmd(templates["cmd"], block["lines"])
        elif kind in ("h1", "head", "body", "cap"):
            element = mk_text(templates[kind], block["text"])
        else:
            raise SystemExit(f"未知的 block type: {kind!r}")
        cell._tc.append(element)
        count[kind] += 1

    if job.get("date"):
        set_cell_text(table.rows[2].cells[3], job["date"])
    if job.get("title"):
        set_cell_text(table.rows[3].cells[1], job["title"])
    if job.get("summary"):
        set_cell_text(table.rows[5].cells[1], job["summary"])
    if job.get("row4_min_twips"):
        set_row_min_height(table, 4, int(job["row4_min_twips"]))

    document.save(out)
    print(
        "saved:", out,
        "| blocks:", ", ".join(f"{k}={v}" for k, v in count.items()),
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--job", required=True, help="UTF-8 JSON job file")
    args = parser.parse_args()
    with open(args.job, encoding="utf-8") as handle:
        job = json.load(handle)
    build(job)


if __name__ == "__main__":
    main()
